"""Validate the final submission package and write a file inventory."""
from pathlib import Path
import json
import re
import sys

import nbformat
import pandas as pd
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
required = [
    "ReadmitRisk_Notebook.ipynb", "ReadmitRisk_Final_Presentation.pptx", "app.py",
    "requirements.txt", "requirements.lock.txt", "README.md", "artifacts/metrics.json",
    "artifacts/classifier.joblib", "artifacts/base_pipeline.joblib", "artifacts/gru.keras",
    "artifacts/lstm.keras", "artifacts/clustering.joblib", "artifacts/model_metadata.json",
    "docs/CRITERIA_CHECKLIST.md", "docs/VIVA_QUESTIONS.md", "docs/FINAL_REPORT.md",
    "figures/dashboard.png", "tables/classifier_metrics.csv", "tables/sequence_model_comparison.csv",
]
for name in required:
    assert (ROOT / name).is_file() and (ROOT / name).stat().st_size > 0, name
notebook = nbformat.read(ROOT / "ReadmitRisk_Notebook.ipynb", as_version=4)
nbformat.validate(notebook)
cells = [c for c in notebook.cells if c.cell_type == "code"]
assert all(c.execution_count is not None for c in cells), "Unexecuted notebook cells"
errors = [o for c in cells for o in c.outputs if o.output_type == "error"]
assert not errors, errors
for cell in cells:
    compile(cell.source, "notebook_cell", "exec")
presentation = Presentation(ROOT / "ReadmitRisk_Final_Presentation.pptx")
assert 12 <= len(presentation.slides) <= 14
summary = json.loads((ROOT / "artifacts/metrics.json").read_text())
assert all(key in summary for key in ["data", "eda", "classifier", "sequences", "allocation", "clustering", "feature_selection", "fairness"])
assert summary["eda"]["chart_count"] >= 12
test_log = (ROOT / "artifacts/test_run.log").read_text()
match = re.search(r"(\d+) passed", test_log)
assert match, "No passing test result captured"
result = {
    "python": sys.version.split()[0], "tests_passed": int(match.group(1)),
    "notebook_code_cells_executed": len(cells), "notebook_errors": len(errors),
    "notebook_inline_figures": sum("image/png" in o.get("data", {}) for c in cells for o in c.outputs),
    "presentation_slides": len(presentation.slides), "eda_figures": summary["eda"]["chart_count"],
    "png_figures": len(list((ROOT / "figures").glob("*.png"))),
    "csv_tables": len(list((ROOT / "tables").glob("*.csv"))),
    "classifier_accuracy_requirement_met": summary["classifier"]["accuracy_requirement_met"],
    "patient_overlap": summary["data"]["patient_overlap"],
    "budget_feasible": summary["allocation"]["optimized"]["budget_feasible"],
    "default_budget_proven_optimal": summary["allocation"]["optimized"]["proven_optimal"],
}
(ROOT / "artifacts/validation.json").write_text(json.dumps(result, indent=2))
inventory = []
for path in sorted(ROOT.rglob("*")):
    relative = path.relative_to(ROOT)
    if not path.is_file() or (relative.parts[0] in {".venv", ".git", ".agents", ".codex", ".pytest_cache"} or "__pycache__" in relative.parts):
        continue
    if path.suffix == ".pyc" or str(relative) in {"diabetic_data.csv", "SUBMISSION_FILES.csv"}:
        continue
    inventory.append({"file": str(relative), "bytes": path.stat().st_size})
pd.DataFrame(inventory).to_csv(ROOT / "SUBMISSION_FILES.csv", index=False)
print(json.dumps(result, indent=2))
